from collections.abc import Callable, Iterator
from datetime import date, datetime
from hashlib import sha256
from pathlib import Path
from typing import Any
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from soybean_sim.data.world_bank import WorldBankDataError, load_world_bank_prices
from soybean_sim.models import Commodity


@pytest.fixture
def sample_path() -> Path:
    """Locate the small extract copied from the verified local workbook."""
    return (
        Path(__file__).resolve().parents[2]
        / "fixtures/world_bank/pink_sheet_sample.xlsx"
    )


@pytest.fixture
def workbook(sample_path: Path) -> Iterator[Workbook]:
    """Provide an editable fixture copy and close its resources after each test."""
    book = load_workbook(sample_path)
    try:
        yield book
    finally:
        book.close()


@pytest.fixture
def save_workbook(workbook: Workbook, tmp_path: Path) -> Callable[[], Path]:
    """Save modified workbooks to a test-local path without touching the fixture."""

    def save() -> Path:
        """Write the modified workbook and return its path."""
        path = tmp_path / "prices.xlsx"
        workbook.save(path)
        return path

    return save


def test_public_api_and_verified_prices(sample_path: Path) -> None:
    """Load all three verified series with exact source values and monthly dates."""
    prices = load_world_bank_prices(str(sample_path))
    expected = [
        (date(1960, 1, 1), (94, 204, 92)),
        (date(1960, 2, 1), (91, 201, 87)),
        (date(1960, 3, 1), (92, 201, 84)),
        (date(2026, 7, 1), (478, 1658, 410)),
        (date(2026, 8, 1), (482, 1639, 406)),
        (date(2026, 9, 1), (524, 1615, 447)),
    ]
    series = (Commodity.SOYBEANS, Commodity.SOYBEAN_OIL, Commodity.SOYBEAN_MEAL)
    assert isinstance(prices, tuple)
    assert len(prices) == 18
    assert {
        (item.month, item.commodity): item.price_usd_per_tonne for item in prices
    } == {
        (month, commodity): value
        for month, values in expected
        for commodity, value in zip(series, values, strict=True)
    }
    assert {item.data_source for item in prices} == {"World Bank Pink Sheet"}
    assert issubclass(WorldBankDataError, ValueError)
    assert all(item.month.day == 1 for item in prices)


@pytest.mark.parametrize(
    ("filters", "count"),
    [
        ({"start_month": date(2026, 8, 1)}, 6),
        ({"end_month": date(1960, 2, 1)}, 6),
        ({"commodities": [Commodity.SOYBEAN_MEAL]}, 6),
        ({"commodities": ["soybeans"]}, 6),
        ({"commodities": []}, 0),
        ({"start_month": date(2030, 1, 1)}, 0),
        ({"end_month": date(1900, 1, 1)}, 0),
        ({"start_month": date(2020, 1, 1), "end_month": date(2020, 1, 1)}, 0),
        (
            {
                "start_month": date(2026, 7, 1),
                "end_month": date(2026, 8, 1),
                "commodities": [Commodity.SOYBEAN_OIL],
            },
            2,
        ),
    ],
)
def test_filters_are_inclusive(
    sample_path: Path, filters: dict[str, Any], count: int
) -> None:
    """Select inclusive dates and commodities without generating absent months."""
    expected = load_world_bank_prices(sample_path)
    start = filters.get("start_month", date.min)
    end = filters.get("end_month", date.max)
    selected = set(filters.get("commodities", list(Commodity)))
    actual = load_world_bank_prices(sample_path, **filters)
    assert len(actual) == count
    assert actual == tuple(
        price
        for price in expected
        if start <= price.month <= end and price.commodity in selected
    )


@pytest.mark.parametrize(
    "filters",
    [
        {"start_month": "2026-01-01"},
        {"start_month": datetime(2026, 1, 1)},
        {"start_month": date(2026, 1, 2)},
        {"end_month": True},
        {"end_month": date(2026, 1, 2)},
        {"start_month": date(2026, 2, 1), "end_month": date(2026, 1, 1)},
        {"commodities": ["unknown"]},
    ],
)
def test_invalid_filters_remain_argument_errors(
    sample_path: Path, filters: dict[str, Any]
) -> None:
    """Keep invalid caller arguments distinct from invalid source workbooks."""
    with pytest.raises(ValueError) as caught:
        load_world_bank_prices(sample_path, **filters)
    assert not isinstance(caught.value, WorldBankDataError)


def test_column_discovery_and_ordering(
    sample_path: Path, workbook: Workbook, save_workbook: Callable[[], Path]
) -> None:
    """Discover moved columns and sort reversed source rows deterministically."""
    sheet = workbook["Monthly Prices"]
    sheet.insert_cols(2)
    rows = list(sheet.iter_rows(min_row=7, values_only=True))
    sheet.delete_rows(7, len(rows))
    for row in reversed(rows):
        sheet.append(row)
    actual = load_world_bank_prices(save_workbook())
    assert actual == load_world_bank_prices(sample_path)
    assert [(item.month, item.commodity.value) for item in actual] == sorted(
        (item.month, item.commodity.value) for item in actual
    )


@pytest.mark.parametrize("missing", [None, "", " ", "...", "…", " … "])
def test_missing_prices_preserve_other_series(
    workbook: Workbook, save_workbook: Callable[[], Path], missing: Any
) -> None:
    """Skip only missing individual prices, leaving the month's other series intact."""
    workbook["Monthly Prices"]["Y7"] = missing
    prices = load_world_bank_prices(save_workbook())
    assert len(prices) == 17
    assert {item.commodity for item in prices if item.month == date(1960, 1, 1)} == {
        Commodity.SOYBEAN_OIL,
        Commodity.SOYBEAN_MEAL,
    }


@pytest.mark.parametrize(
    "invalid",
    [
        True,
        False,
        "bad",
        "NA",
        "N/A",
        "#N/A",
        "NaN",
        "inf",
        "-inf",
        0,
        -1,
        "0",
        "-1",
        "1_000",
        "1,000",
        "1e9999",
        "1e-9999",
        "=1+1",
        date(2026, 1, 1),
    ],
)
def test_invalid_prices_have_observation_context(
    workbook: Workbook, save_workbook: Callable[[], Path], invalid: Any
) -> None:
    """Identify the file, sheet, row, month, and commodity of invalid prices."""
    workbook["Monthly Prices"]["Y7"] = invalid
    path = save_workbook()
    with pytest.raises(WorldBankDataError) as caught:
        load_world_bank_prices(path)
    message = str(caught.value)
    for context in (str(path), "Monthly Prices", "row 7", "soybeans", "1960-01"):
        assert context in message


@pytest.mark.parametrize(
    "invalid",
    [
        None,
        202609,
        "2026M00",
        "2026M13",
        "0000M01",
        "2026M9",
        "2026m09",
        "2026-09",
        date(2026, 9, 1),
    ],
)
def test_invalid_months_are_not_ignored(
    workbook: Workbook, save_workbook: Callable[[], Path], invalid: Any
) -> None:
    """Validate month labels even when a date filter would exclude valid early rows."""
    workbook["Monthly Prices"]["A7"] = invalid
    with pytest.raises(WorldBankDataError, match=r"row 7.*invalid month"):
        load_world_bank_prices(save_workbook(), start_month=date(2026, 1, 1))


@pytest.mark.parametrize("duplicate_value", [94, "94.000", "9.4e1", None])
def test_equivalent_or_missing_duplicates_do_not_add_prices(
    workbook: Workbook, save_workbook: Callable[[], Path], duplicate_value: Any
) -> None:
    """Count equivalent known prices once and skip missing duplicate cells."""
    sheet = workbook["Monthly Prices"]
    row = [cell.value for cell in sheet[7]]
    row[24] = duplicate_value
    sheet.append(row)
    assert len(load_world_bank_prices(save_workbook())) == 18


@pytest.mark.parametrize("conflict", [95, "94.000000000000000000001"])
def test_conflicting_duplicates_preserve_decimal_precision(
    workbook: Workbook, save_workbook: Callable[[], Path], conflict: Any
) -> None:
    """Detect conflicting source values even when they round to the same float."""
    sheet = workbook["Monthly Prices"]
    row = [cell.value for cell in sheet[7]]
    row[24] = conflict
    sheet.append(row)
    with pytest.raises(
        WorldBankDataError, match=r"row 13.*soybeans.*conflicting duplicate"
    ):
        load_world_bank_prices(save_workbook())


@pytest.mark.parametrize("column", ["Y", "Z", "AA"])
def test_all_series_are_required(
    workbook: Workbook, save_workbook: Callable[[], Path], column: str
) -> None:
    """Validate complete supported headers before applying commodity filters."""
    workbook["Monthly Prices"][f"{column}5"] = None
    with pytest.raises(WorldBankDataError, match="missing commodity columns"):
        load_world_bank_prices(save_workbook(), commodities=[])


@pytest.mark.parametrize("unit", [None, "USD/mt", "($/kg)", 1000])
def test_units_are_verified_for_every_series(
    workbook: Workbook, save_workbook: Callable[[], Path], unit: Any
) -> None:
    """Require the verified source label even for an unselected commodity."""
    workbook["Monthly Prices"]["Z6"] = unit
    with pytest.raises(
        WorldBankDataError, match=r"row 6, Soybean oil.*unexpected unit"
    ):
        load_world_bank_prices(save_workbook(), commodities=[Commodity.SOYBEANS])


@pytest.mark.parametrize("column", ["A", "AB"])
def test_duplicate_or_misplaced_headers_are_rejected(
    workbook: Workbook, save_workbook: Callable[[], Path], column: str
) -> None:
    """Keep each series unique and separate from the fixed month column."""
    workbook["Monthly Prices"][f"{column}5"] = "Soybeans"
    with pytest.raises(WorldBankDataError, match=r"row 5.*duplicate or misplaced"):
        load_world_bank_prices(save_workbook())


@pytest.mark.parametrize("structure", ["missing_sheet", "empty_sheet", "missing_units"])
def test_invalid_workbook_structure(
    workbook: Workbook, save_workbook: Callable[[], Path], structure: str
) -> None:
    """Report missing worksheets and incomplete header layouts explicitly."""
    sheet = workbook["Monthly Prices"]
    if structure == "missing_sheet":
        workbook.remove(sheet)
    elif structure == "empty_sheet":
        sheet.delete_rows(1, sheet.max_row)
    else:
        sheet.delete_rows(6, sheet.max_row)
    with pytest.raises(WorldBankDataError, match="Monthly Prices"):
        load_world_bank_prices(save_workbook())


def test_blank_rows_header_only_and_unrelated_sheets(
    workbook: Workbook, save_workbook: Callable[[], Path]
) -> None:
    """Ignore blank rows and unrelated sheets; accept a valid header-only dataset."""
    sheet = workbook["Monthly Prices"]
    sheet.insert_rows(9)
    workbook["Description"]["A2"] = "not a month"
    assert len(load_world_bank_prices(save_workbook())) == 18
    sheet.delete_rows(7, sheet.max_row)
    assert load_world_bank_prices(save_workbook()) == ()


def test_filtering_precedes_price_validation(
    workbook: Workbook, save_workbook: Callable[[], Path]
) -> None:
    """Skip invalid prices in excluded months and unselected commodities."""
    sheet = workbook["Monthly Prices"]
    sheet["Y7"] = "bad"
    sheet["Z10"] = "bad"
    prices = load_world_bank_prices(
        save_workbook(), start_month=date(2026, 1, 1), commodities=[Commodity.SOYBEANS]
    )
    assert len(prices) == 3


@pytest.mark.parametrize(
    ("changes", "expected_error"),
    [
        (
            {"Y6": "($/kg)", "Z5": "Soybeans"},
            r"row 6, Soybeans.*unexpected unit",
        ),
        (
            {"Z5": "Soybeans", "Z6": "($/kg)"},
            r"row 5.*duplicate or misplaced",
        ),
        (
            {"AA5": None, "A7": "not a month"},
            r"row 5.*missing commodity columns",
        ),
        (
            {"Z6": "($/kg)", "A7": "not a month"},
            r"row 6, Soybean oil.*unexpected unit",
        ),
    ],
)
def test_schema_validation_preserves_source_order(
    workbook: Workbook,
    save_workbook: Callable[[], Path],
    changes: dict[str, Any],
    expected_error: str,
) -> None:
    """Preserve error priority when multiple header, unit, or row problems coexist."""
    sheet = workbook["Monthly Prices"]
    for address, cell_value in changes.items():
        sheet[address] = cell_value
    with pytest.raises(WorldBankDataError, match=expected_error):
        load_world_bank_prices(save_workbook(), commodities=[])


def test_empty_commodity_selection_still_validates_months(
    workbook: Workbook, save_workbook: Callable[[], Path]
) -> None:
    """Validate nonblank month labels even when no price series are selected."""
    workbook["Monthly Prices"]["A7"] = "not a month"
    with pytest.raises(WorldBankDataError, match=r"row 7.*invalid month"):
        load_world_bank_prices(save_workbook(), commodities=[])


@pytest.mark.parametrize("kind", ["missing", "directory", "not_xlsx", "bad_xml"])
def test_unreadable_sources_have_path_and_cause(tmp_path: Path, kind: str) -> None:
    """Translate file and container failures without losing their original cause."""
    path = tmp_path / "prices.xlsx"
    if kind == "directory":
        path.mkdir()
    elif kind == "not_xlsx":
        path.write_text("not an Excel file", encoding="utf-8")
    elif kind == "bad_xml":
        with ZipFile(path, "w") as archive:
            archive.writestr("[Content_Types].xml", "<unclosed")
    with pytest.raises(WorldBankDataError) as caught:
        load_world_bank_prices(path)
    assert str(path) in str(caught.value)
    assert caught.value.__cause__ is not None


@pytest.mark.parametrize("number", ["1e309", "-1e309"])
def test_numeric_overflow_cells_are_not_missing(
    sample_path: Path, tmp_path: Path, number: str
) -> None:
    """Reject overflowing numeric cells instead of treating them as missing."""
    path = tmp_path / "overflow.xlsx"
    namespace = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(sample_path) as source, ZipFile(path, "w") as target:
        for name in source.namelist():
            content = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(content)
                cell = root.find(".//m:c[@r='Y7']/m:v", namespace)
                assert cell is not None
                cell.text = number
                content = ElementTree.tostring(root)
            target.writestr(name, content)
    with pytest.raises(WorldBankDataError, match=r"row 7.*finite and positive"):
        load_world_bank_prices(path)


def test_source_file_is_unchanged(sample_path: Path) -> None:
    """Leave the XLSX bytes unchanged after loading."""
    before = sha256(sample_path.read_bytes()).digest()
    load_world_bank_prices(sample_path)
    assert sha256(sample_path.read_bytes()).digest() == before
