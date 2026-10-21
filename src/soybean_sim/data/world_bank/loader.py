"""Read local Pink Sheet XLSX observations into validated global reference prices."""

from collections.abc import Iterable, Iterator, Mapping
from contextlib import contextmanager
from datetime import date
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from soybean_sim.models import Commodity, CommodityPrice

from . import mappings, parsing
from .parsing import WorldBankDataError

_PriceKey = tuple[date, Commodity]
_Observations = dict[_PriceKey, parsing.ParsedPrice]
_SelectedRow = tuple[date, tuple[object, ...], str]


def load_world_bank_prices(
    path: str | Path,
    *,
    start_month: date | None = None,
    end_month: date | None = None,
    commodities: Iterable[Commodity] | None = None,
) -> tuple[CommodityPrice, ...]:
    """Load local XLSX prices sorted by month and commodity, without changing the file.

    Bounds are inclusive first-of-month dates. None selects all commodities; an
    empty selection or no matching observations returns an empty tuple. Blanks
    and ellipsis markers are skipped. Equivalent duplicates count once, while
    conflicting observations are invalid. No conversion or interpolation occurs.
    """
    parsing.validate_month_range(start_month, end_month)
    selected_commodities = (
        set(Commodity)
        if commodities is None
        else {Commodity(commodity) for commodity in commodities}
    )

    workbook_path = Path(path)
    with _open_workbook(workbook_path) as workbook:
        worksheet = _price_sheet(workbook, workbook_path)
        return _read_prices(
            worksheet, workbook_path, start_month, end_month, selected_commodities
        )


@contextmanager
def _open_workbook(workbook_path: Path) -> Iterator[Workbook]:
    """Manage workbook resources and translate file/read failures during ingestion."""
    try:
        with workbook_path.open("rb") as stream:
            workbook = load_workbook(stream, data_only=False, keep_links=False)
            try:
                yield workbook
            finally:
                workbook.close()
    except WorldBankDataError:
        raise
    except (
        OSError,
        BadZipFile,
        InvalidFileException,
        ParseError,
        KeyError,
        ValueError,
    ) as exc:
        raise WorldBankDataError(
            f"{workbook_path}: cannot read Pink Sheet workbook: {exc}"
        ) from exc


def _price_sheet(workbook: Workbook, path: Path) -> Worksheet:
    """Locate the monthly data sheet and check the minimum header layout."""
    if mappings.SHEET_NAME not in workbook.sheetnames:
        raise WorldBankDataError(f"{path}: missing worksheet {mappings.SHEET_NAME!r}")

    worksheet = workbook[mappings.SHEET_NAME]
    if not isinstance(worksheet, Worksheet) or worksheet.max_row < mappings.UNIT_ROW:
        raise WorldBankDataError(
            f"{path}: worksheet {mappings.SHEET_NAME!r}: missing header or unit rows"
        )

    return worksheet


def _series_columns(worksheet: Worksheet, context: str) -> dict[Commodity, int]:
    """Find all three series by name and validate units before filtering."""
    columns = {}
    for column_number, header in _iter_series_headers(worksheet):
        commodity = mappings.SERIES[header]
        if commodity in columns or column_number == mappings.MONTH_COLUMN:
            raise WorldBankDataError(
                f"{context}: row {mappings.HEADER_ROW}: "
                f"duplicate or misplaced {header!r}"
            )

        source_unit = worksheet.cell(row=mappings.UNIT_ROW, column=column_number).value
        _validate_price_unit(source_unit, header, context)
        columns[commodity] = column_number - 1
    _check_required_series(columns, context)

    return columns


def _iter_series_headers(worksheet: Worksheet) -> Iterator[tuple[int, str]]:
    """Yield column numbers and trimmed names for supported commodity headers."""
    for column_number, cell in enumerate(worksheet[mappings.HEADER_ROW], start=1):
        if not isinstance(cell.value, str):
            continue

        header = cell.value.strip()
        if header in mappings.SERIES:
            yield column_number, header


def _check_required_series(columns: Mapping[Commodity, int], context: str) -> None:
    """Check that the discovered columns cover every required commodity series."""
    missing = set(mappings.SERIES.values()) - columns.keys()
    if missing:
        raise WorldBankDataError(
            f"{context}: row {mappings.HEADER_ROW}: missing commodity columns: "
            f"{', '.join(sorted(commodity.value for commodity in missing))}"
        )


def _validate_price_unit(source_unit: object, series_name: str, context: str) -> None:
    """Check a named series uses the verified USD-per-tonne unit label."""
    if not isinstance(source_unit, str) or source_unit.strip() != mappings.PRICE_UNIT:
        raise WorldBankDataError(
            f"{context}: row {mappings.UNIT_ROW}, {series_name}: "
            f"unexpected unit {source_unit!r}; expected {mappings.PRICE_UNIT!r}"
        )


def _read_prices(
    worksheet: Worksheet,
    path: Path,
    start_month: date | None,
    end_month: date | None,
    selected_commodities: set[Commodity],
) -> tuple[CommodityPrice, ...]:
    """Orchestrate schema validation, row selection, collection, and model creation."""
    context = f"{path}: worksheet {mappings.SHEET_NAME!r}"
    columns = _series_columns(worksheet, context)
    selected_columns = {
        commodity: column
        for commodity, column in columns.items()
        if commodity in selected_commodities
    }
    rows = _selected_rows(worksheet, (start_month, end_month), context)
    observations = _collect_observations(rows, selected_columns)

    return _build_prices(observations)


def _selected_rows(
    worksheet: Worksheet,
    month_bounds: tuple[date | None, date | None],
    context: str,
) -> Iterator[_SelectedRow]:
    """Yield month, cells, and context after blank-row checks and month selection."""
    start_month, end_month = month_bounds
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=mappings.FIRST_DATA_ROW, values_only=True),
        start=mappings.FIRST_DATA_ROW,
    ):
        if all(
            cell is None or (isinstance(cell, str) and not cell.strip()) for cell in row
        ):
            continue

        row_context = f"{context}: row {row_number}"
        month = parsing.parse_month(row[mappings.MONTH_COLUMN - 1], row_context)
        if start_month is not None and month < start_month:
            continue

        if end_month is not None and month > end_month:
            continue

        yield month, row, row_context


def _collect_observations(
    rows: Iterable[_SelectedRow], columns: dict[Commodity, int]
) -> _Observations:
    """Parse selected series, skip missing cells, and enforce duplicate consistency."""
    observations: _Observations = {}
    for month, row, row_context in rows:
        for commodity, column in columns.items():
            observation_context = f"{row_context}, {commodity.value}, {month:%Y-%m}"
            price = parsing.parse_price(row[column], observation_context)
            if price is not None:
                _record_price(
                    observations, (month, commodity), price, observation_context
                )

    return observations


def _build_prices(observations: _Observations) -> tuple[CommodityPrice, ...]:
    """Construct immutable price models sorted by month and commodity identifier."""
    return tuple(
        CommodityPrice(
            month=month,
            commodity=commodity,
            price_usd_per_tonne=price.usd_per_tonne,
            data_source=mappings.DATA_SOURCE,
        )
        for (month, commodity), price in sorted(observations.items())
    )


def _record_price(
    observations: _Observations,
    key: _PriceKey,
    price: parsing.ParsedPrice,
    context: str,
) -> None:
    """Retain equivalent duplicates once, comparing decimals before float conversion."""
    previous = observations.get(key)
    if previous is not None and previous.source_value != price.source_value:
        raise WorldBankDataError(f"{context}: conflicting duplicate price")
    observations[key] = price
